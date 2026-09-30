"""Fixes from the code review: dated FX, export columns, Excel safety, stale reviews,
full review prompts and product groups."""
import io
import json

import pandas as pd
from openpyxl import load_workbook

from price_lens.core import fx
from price_lens.orchestrator import price_report

ECB_XML = """<?xml version="1.0" encoding="UTF-8"?>
<gesmes:Envelope><Cube><Cube time='2026-09-28'>
<Cube currency='USD' rate='1.1500'/><Cube currency='GBP' rate='0.8500'/>
<Cube currency='PLN' rate='4.2500'/></Cube></Cube></gesmes:Envelope>"""


# --- dated exchange rates ----------------------------------------------------
def test_parse_ecb_converts_to_usd_per_unit():
    day, rates = fx.parse_ecb(ECB_XML)
    assert day == "2026-09-28"
    assert rates["EUR"] == 1.15 and rates["USD"] == 1.0
    assert abs(rates["GBP"] - 1.15 / 0.85) < 1e-6
    assert abs(rates["PLN"] - 1.15 / 4.25) < 1e-6


def test_get_rates_live_then_cached_then_static(tmp_path):
    live = fx.get_rates(tmp_path, fetch=lambda: ECB_XML)
    assert live["date"] == "2026-09-28" and "European Central Bank" in live["source"]
    assert "EUR" not in live["static_currencies"]
    assert (tmp_path / fx.CACHE_FILE).exists()

    def offline():
        raise OSError("no network")

    cached = fx.get_rates(tmp_path, fetch=offline)
    assert cached["date"] == "2026-09-28" and "cached" in cached["source"]

    static = fx.get_rates(tmp_path / "empty", fetch=offline)
    assert static["date"] is None and "Static" in static["source"]
    assert "static" in fx.describe(static).lower()


def test_fx_saved_per_run_and_applied(tmp_path):
    info = fx.get_rates(tmp_path, fetch=lambda: ECB_XML)
    fx.save_for_run(tmp_path, info)
    assert fx.load_for_run(tmp_path)["date"] == "2026-09-28"
    assert "2026-09-28" in fx.describe(fx.load_for_run(tmp_path))
    frame = pd.DataFrame({"price_value": [100, 50, 10], "currency_code": ["EUR", "gbp", "XYZ"],
                          "price_usd": [1.0, 2.0, 3.0]})
    out = fx.apply(frame, info["rates"])
    assert list(out["price_usd"]) == [115.0, round(50 * 1.15 / 0.85, 2), 3.0]  # unknown kept


# --- report / exports ----------------------------------------------------------
def _frame():
    return pd.DataFrame([
        {"search_id": 1, "product_id": "a1", "platform": "amazon", "marketplace": "amazon.de",
         "product_name": "=HYPERLINK(\"http://evil\",\"x\")", "price_value": 100,
         "currency_code": "EUR", "price_usd": 115, "scraped_at": "2026-09-28T10:00:00"},
        {"search_id": 2, "product_id": "b1", "platform": "ebay", "marketplace": "ebay.de",
         "product_name": "Hoodie", "price_value": 40, "currency_code": "EUR", "price_usd": 46,
         "scraped_at": "2026-09-28T10:05:00", "relevance_confidence_pct": "88"},
    ])


def test_prepare_products_without_country_column_and_groups():
    labels = price_report.group_labels({"searches": [{"product_type": "Sneakers"},
                                                     {"query": {"default": "hoodie"}}]})
    assert labels == {1: "1 · Sneakers", 2: "2 · hoodie"}
    df = price_report.prepare_products(_frame(), labels)
    assert list(df["country"]) == ["Germany", "Germany"]
    assert list(df["group"]) == ["1 · Sneakers", "2 · hoodie"]
    assert df["relevance_confidence_pct"].dtype.kind == "f"


def test_exports_keep_ids_timestamp_and_confidence():
    df = price_report.prepare_products(_frame())
    exported = price_report.product_view(df, export=True)
    for column in ("product_id", "search_id", "scraped_at", "relevance_confidence_pct"):
        assert column in exported.columns


def test_excel_report_does_not_run_formulas_and_states_fx(tmp_path):
    df = price_report.prepare_products(_frame())
    summary = price_report.price_summary(df)
    info = fx.get_rates(tmp_path, fetch=lambda: ECB_XML)
    data = price_report.price_report_xlsx(summary, df, {"run": "r", "group": "1 · Sneakers"}, info)
    book = load_workbook(io.BytesIO(data))
    products = book["Products"]
    cells = [c for row in products.iter_rows() for c in row if str(c.value).startswith("=")]
    assert cells and all(c.data_type == "s" for c in cells)  # stored as text, not formula
    about = " ".join(str(c.value) for row in book["About"].iter_rows() for c in row if c.value)
    assert "2026-09-28" in about and "European Central Bank" in about


# --- review: full prompt, stale after retry, add only new listings --------------
def _review(tmp_path, rows):
    from price_lens.agent.review import prepare_review_batches

    source = tmp_path / "detailed_products.csv"
    pd.DataFrame(rows).to_csv(source, index=False)
    prepare_review_batches(source, tmp_path / "ai_review", research_question="Shoes",
                           instruction="Keep shoes", batch_size=5)
    return source, tmp_path / "ai_review"


def _row(i, name, **extra):
    return {"search_id": 1, "platform": "zalando", "marketplace": "zalando.de",
            "product_name": name, "price_value": 50.0 + i, "currency_code": "EUR",
            "product_url": f"https://x/{i}", "product_id": f"id{i}", **extra}


def test_prompt_has_full_title_colour_and_availability(tmp_path):
    from price_lens.agent import assist

    long_title = "Running shoe " + "extra words " * 12 + "in black"
    _, review_dir = _review(tmp_path, [_row(0, long_title, color="black",
                                            availability="in stock")])
    prompt = assist.compact_prompt(list(assist.batch_files(review_dir).values()))
    assert "in black" in prompt and "color: black" in prompt
    assert "availability: in stock" in prompt


def test_retry_marks_review_outdated_and_new_rows_are_added(tmp_path):
    from price_lens.agent import assist
    from price_lens.agent.review import (
        apply_review_decisions,
        extend_review_batches,
        mark_review_outdated,
        review_outdated,
    )

    source, review_dir = _review(tmp_path, [_row(0, "Shoe A"), _row(1, "Shoe B")])
    assist.save_answer(review_dir, assist.pending_batches(review_dir),
                       "p001|keep|0.9|shoe|ok\np002|keep|0.9|shoe|ok")
    apply_review_decisions(review_dir)
    assert assist.pending_batches(review_dir) == []

    # a retry adds one listing
    pd.DataFrame([_row(0, "Shoe A"), _row(1, "Shoe B"), _row(2, "Shoe C")]).to_csv(source, index=False)
    mark_review_outdated(review_dir, "retry retry_1", 1)
    assert review_outdated(review_dir)["added"] == 1

    assert extend_review_batches(review_dir, source)["added"] == 1
    assert review_outdated(review_dir) is None
    assert not (review_dir / "final_products.csv").exists()  # must be applied again
    pending = assist.pending_batches(review_dir)
    assert len(pending) == 1 and pending[0].endswith("_new")
    prompt = assist.compact_prompt([assist.batch_files(review_dir)[pending[0]]])
    assert "Shoe C" in prompt and "Shoe A" not in prompt
    assist.save_answer(review_dir, pending, "p003|drop|0.8|other|not a shoe")
    report = apply_review_decisions(review_dir)
    assert report["accepted"] == 2 and report["excluded"] == 1


def test_merge_retry_flags_existing_review(tmp_path):
    from price_lens.agent.review import review_outdated
    from price_lens.orchestrator.retry import merge_retry

    parent, retry_dir = tmp_path / "analyst_1", tmp_path / "retry_1"
    parent.mkdir()
    retry_dir.mkdir()
    pd.DataFrame([_row(0, "Shoe A")]).to_csv(parent / "detailed_products.csv", index=False)
    pd.DataFrame([_row(0, "Shoe A")]).to_csv(parent / "raw_products.csv", index=False)
    (parent / "collection_report.json").write_text(json.dumps({"searches": []}))
    _review(parent, [_row(0, "Shoe A")])
    pd.DataFrame([_row(5, "Shoe Z")]).to_csv(retry_dir / "detailed_products.csv", index=False)
    (retry_dir / "analyst_plan.json").write_text(json.dumps({"searches": [{"origin_search_id": 1}]}))
    result = merge_retry(parent, retry_dir)
    assert result == {"added": 1, "review_outdated": True}
    assert review_outdated(parent / "ai_review")["reasons"] == ["retry retry_1"]
