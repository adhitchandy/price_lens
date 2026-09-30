import copy
import json
from pathlib import Path
from types import SimpleNamespace

import pandas as pd
import pytest

from price_lens.core.validation import RequestValidationError
from price_lens.orchestrator.analyst import (
    compile_analyst_plan,
    load_analyst_text,
    run_analyst_plan,
)

ROOT = Path(__file__).parents[1]


def example():
    return json.loads((ROOT / "examples/analyst_plan.json").read_text())


def test_mixed_country_scope_and_localized_queries():
    plans, preview = compile_analyst_plan(example())
    assert len(plans) == 2
    assert len(plans[0].marketplaces["zalando"]) == 28
    assert plans[1].marketplaces["mediamarkt"] == ("mediamarkt.de", "mediamarkt.at")
    assert plans[1].marketplaces["amazon"] == ("amazon.de",)
    assert {r["country"] for r in preview if r["search_id"] == 2} == {"Germany", "Austria"}
    assert all(r["query"] == "Waschmaschine" for r in preview if r["search_id"] == 2)


def test_text_fence_import_and_missing_translation():
    data = load_analyst_text("\ufeff```json\n" + json.dumps(example()) + "\n```")
    del data["searches"][0]["translations"]["ja"]
    with pytest.raises(RequestValidationError, match="ja translation"):
        compile_analyst_plan(data)


@pytest.mark.parametrize("value", ["Germany", ["Atlantis"]])
def test_bad_countries_fail_on_import(value):
    data = example()
    data["searches"][0]["countries"] = value
    with pytest.raises(RequestValidationError):
        load_analyst_text(json.dumps(data))


def test_website_changes_do_not_broaden_other_searches():
    data = example()
    data["searches"][0]["platforms"] = ["ebay"]
    plans, _ = compile_analyst_plan(data)
    assert plans[0].platforms == ("ebay",)
    assert plans[1].platforms == ("amazon", "ebay", "mediamarkt")


def test_export_preserves_partial_result_and_does_not_invent_confidence(tmp_path):
    calls = []

    def runner(plan, log):
        calls.append(plan)
        if len(calls) == 2:
            raise RuntimeError("Browser closed")
        rows = pd.DataFrame([{"product_name": "=malicious()", "currency_code": "EUR",
                              "price_value": 89.95, "platform": "ebay", "raw_extra": "audit"}])
        return SimpleNamespace(raw_products=rows, cleaned_products=rows,
                               report={"status": "completed"}, run_dir=tmp_path)

    root, final, reports = run_analyst_plan(example(), tmp_path, runner=runner)
    assert len(final) == 1
    assert final.relevance_confidence_pct.isna().all()
    assert final.review_status.tolist() == ["not_reviewed"]
    assert reports[1]["status"] == "failed"
    assert "raw_extra" in pd.read_csv(root / "detailed_products.csv")
    assert "raw_extra" not in pd.read_csv(root / "final_products.csv")
    from openpyxl import load_workbook
    sheet = load_workbook(root / "final_products.xlsx").active
    col = [cell.value for cell in sheet[1]].index("product_name") + 1
    assert sheet.cell(2, col).data_type == "s"


def fake_outcome(domain, term):
    from price_lens.core.results import ScrapeOutcome
    from price_lens.core.schemas import ScrapeReport

    rows = pd.DataFrame([
        {"platform": "mediamarkt", "marketplace": domain, "product_name": name,
         "search_term": term, "price_value": price, "price_usd": price, "currency_code": "EUR",
         "product_id": name, "sponsored": False, "rating": 4.0, "origin_country": "Germany"}
        for name, price in (("Phone A", 150.0), ("Phone B", 900.0))
    ])
    report = ScrapeReport(platform="mediamarkt", status="completed", pages_succeeded=1)
    report.events.append({"marketplace": domain, "page": 1, "listings": 2, "status": "succeeded",
                          "final_url": f"https://www.{domain}/"})
    return ScrapeOutcome(rows, report)


def _app():
    from streamlit.testing.v1 import AppTest
    return AppTest.from_file(str(ROOT / "apps/unified_app.py"), default_timeout=60).run()


def _run_button(app):
    return next(b for b in app.button if b.label == "🚀 Run scraper")


def test_unified_ui_load_and_edit_without_browser():
    app = _app()
    assert not app.exception
    assert not _run_button(app).disabled
    data = example()
    data["searches"] = [copy.deepcopy(data["searches"][1])]
    next(t for t in app.text_area if t.label == "…or paste JSON").set_value(json.dumps(data))
    next(b for b in app.button if b.label == "Load Plan into Workspace").click().run()
    assert not app.exception
    plan = app.session_state["plan"]
    assert plan["schema_version"] == "analyst-v2"
    assert [t["platform"] for t in plan["searches"][0]["targets"]] == ["amazon", "ebay", "mediamarkt"]
    # imported values must not be overwritten by stale widget state
    assert next(w for w in app.text_input if w.label == "Main query term").value == "washing machines"
    assert not _run_button(app).disabled


def test_unified_ui_run_goes_through_background_job(monkeypatch, tmp_path):
    from price_lens.orchestrator import jobs, registry
    from price_lens.orchestrator.registry import MarketplaceAdapter

    seen = []

    def collect(plan, _log):
        domain = plan.marketplaces[0]
        seen.append((domain, plan.searches[0].search_term, plan.searches[0].rule.exclude,
                     plan.max_products))
        return fake_outcome(domain, plan.searches[0].search_term)

    def start_inline(plan, output_root, kind="scrape"):
        run_dir = output_root / f"analyst_test_{kind}"
        run_dir.mkdir(parents=True)
        (run_dir / "analyst_plan.json").write_text(json.dumps(plan), encoding="utf-8")
        jobs._update(run_dir, kind=kind, status="queued", heartbeat=jobs._now())
        jobs.run_worker(run_dir)
        return run_dir

    monkeypatch.setenv("PI_OUTPUT_DIR", str(tmp_path))
    monkeypatch.setitem(registry.MARKETPLACE_REGISTRY, "mediamarkt", MarketplaceAdapter("mediamarkt", collect))
    monkeypatch.setattr(jobs, "start_job", start_inline)
    app = _app()
    app.session_state["plan"]["searches"][0]["query"]["filters"]["max_price"] = 500.0
    app.session_state["rev"] = 99
    app.run()
    _run_button(app).click().run()
    assert not app.exception
    app.run()  # monitor sees the finished job and loads the result
    result = app.session_state["result"]
    assert result["status"] == "completed"
    assert seen == [("mediamarkt.de", "Smartphone", ("case", "cover", "hülle"), 100)]
    assert list(result["final"]["product_name"]) == ["Phone A"]
    history_rows = [r for r in app.dataframe]
    assert history_rows  # results + history tables render


def test_v1_plan_upgrades_to_equivalent_v2():
    from price_lens.orchestrator.analyst import upgrade_to_v2
    v1 = example()
    v2 = upgrade_to_v2(v1)
    _, preview_v1 = compile_analyst_plan(v1)
    _, preview_v2 = compile_analyst_plan(v2)
    key = lambda rows: sorted((r["search_id"], r["domain"], r["query"], r["exclude"]) for r in rows)  # noqa: E731
    assert key(preview_v1) == key(preview_v2)
