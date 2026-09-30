"""The web app (price_lens.webapp): API functions and the HTTP server."""
import json
import threading
import urllib.error
import urllib.request

import pandas as pd
import pytest

from price_lens.agent import assist
from price_lens.webapp import api, server

PLAN = {
    "schema_version": "analyst-v2", "research_question": "Running shoes",
    "execution": {"products_per_storefront": 50, "retries": 0, "headless": True},
    "searches": [{"id": "s1", "product_type": "Running shoes",
                  "query": {"default": "running shoes", "translations": {"de": "Laufschuhe"}, "filters": {}},
                  "targets": [{"platform": "zalando", "countries": ["Germany", "Poland"]},
                              {"platform": "amazon", "countries": ["Germany"]}]}],
}


def make_run(root, name="analyst_20260929_160000", rows=None):
    run = root / name
    run.mkdir(parents=True)
    rows = rows or [
        {"search_id": 1, "product_id": f"p{i}", "platform": "zalando", "marketplace": m, "product_name": f"Shoe {i}",
         "brand": "Nike", "price_value": 100 + i, "currency_code": "EUR", "price_usd": 116.7 + i,
         "product_url": f"https://www.{m}/p/{i}", "scraped_at": "2026-09-29T16:00:00"}
        for i, m in enumerate(["zalando.de", "zalando.de", "zalando.pl", "zalando.pl", "amazon.de"])]
    frame = pd.DataFrame(rows)
    frame.to_csv(run / "final_products.csv", index=False)
    frame.to_csv(run / "detailed_products.csv", index=False)
    (run / "analyst_plan.json").write_text(json.dumps(PLAN), encoding="utf-8")
    (run / "collection_report.json").write_text(json.dumps({"searches": [
        {"search_id": 1, "status": "completed", "report": {}}]}), encoding="utf-8")
    return run


@pytest.fixture()
def out(tmp_path, monkeypatch):
    monkeypatch.setenv("PI_OUTPUT_DIR", str(tmp_path))
    return tmp_path


def test_drafts_validate_and_list(out):
    draft = api.create_draft({"plan": PLAN, "goal": "Compare shoe prices"})
    assert api.get_draft(draft["id"])["goal"] == "Compare shoe prices"
    check = api.validate_plan(draft["plan"])
    assert check["ok"] and check["storefronts"] == 3 and check["estimate_minutes"] >= 1
    languages = {lang["lang"]: lang for lang in check["searches"][0]["languages"]}
    assert languages["de"]["term"] == "Laufschuhe" and not languages["de"]["needed"]
    assert languages["pl"]["needed"]  # Poland has no local word yet
    saved = api.save_draft(draft["id"], {"goal": "Updated"})
    assert saved["goal"] == "Updated"
    items = api.research_list()["items"]
    assert items[0]["kind"] == "draft" and items[0]["storefronts"] == 3
    api.delete_draft(draft["id"])
    assert api.research_list()["items"] == []
    with pytest.raises(api.ApiError):
        api.get_draft("../../etc/passwd")


def test_quick_setup_prompt_import_and_settings(out):
    search = {"query": {"default": "shoes", "filters": {}}, "targets": [{"platform": "zalando", "countries": ["Germany"]}]}
    res = api.apply_quick_setup({"search": search, "preset": "Fashion & shoes", "group": "DACH"})
    assert {t["platform"] for t in res["search"]["targets"]} == {"zalando", "amazon", "ebay"}
    assert "socks" in res["search"]["query"]["filters"]["exclude"]
    assert "Research request" in api.planning_prompt({"request": "Sneakers in Germany"})["prompt"]
    assert api.import_plan({"text": json.dumps(PLAN)})["plan"]["schema_version"] == "analyst-v2"
    with pytest.raises(api.ApiError):
        api.import_plan({"text": "not json"})
    api.save_settings({"execution": {"products_per_storefront": 250}})
    assert api.settings()["execution"]["products_per_storefront"] == 250
    assert api.new_plan()["execution"]["products_per_storefront"] == 250


def test_run_detail_products_and_export(out):
    run = make_run(out)
    item = api.research_list()["items"][0]
    assert item["kind"] == "run" and item["listings"] == 5 and item["median"] == 118.7
    detail = api.run_detail(run.name)
    assert [s["key"] for s in detail["steps"]] == ["plan", "check", "collect", "review", "results"]
    assert detail["collected"] == 5 and detail["status"] == "completed"
    data = api.products(run.name)
    assert data["source"] == "all" and len(data["rows"]) == 5 and data["groups"] == ["1 · Running shoes"]
    xlsx = api.export(run.name, {"kind": "xlsx", "rows": [0, 1], "filters": "country=Germany"})
    assert xlsx.name.endswith(".xlsx") and xlsx.data[:2] == b"PK"
    csv = api.export(run.name, {"kind": "csv", "rows": [0, 1, 2]}).data.decode("utf-8-sig")
    assert len(csv.strip().splitlines()) == 4
    assert api.raw_file(run.name, "final_products.csv").data.startswith(b"search_id")
    with pytest.raises(api.ApiError):
        api.raw_file(run.name, "../analyst_plan.json")
    with pytest.raises(api.ApiError):
        api.run_detail("not_a_run")
    copy = api.create_draft({"from_run": run.name})
    assert copy["plan"]["research_question"].endswith("(copy)")


def test_review_flow_with_hand_changes(out):
    run = make_run(out)
    state = api.review_prepare(run.name, {"instruction": "Keep shoes", "batch_size": 10})
    assert state["prepared"] and state["counts"]["pending"] == 5
    prompt = api.review_prompt(run.name)
    assert prompt["count"] == 5 and "Shoe 0" in prompt["prompt"]
    answer = "\n".join(f"p00{i + 1}|{'keep' if i < 4 else 'drop'}|0.95|shoe|reason {i}" for i in range(5))
    saved = api.review_answer(run.name, {"batches": prompt["batches"], "text": answer})
    assert saved["saved"] == 5 and saved["state"]["counts"] == {"accepted": 4, "uncertain": 0, "excluded": 1, "pending": 0}
    rows = api.review_decisions(run.name)["rows"]
    dropped = next(r for r in rows if r["status"] == "excluded")
    api.review_apply(run.name)
    assert api.products(run.name)["source"] == "reviewed" and len(api.products(run.name)["rows"]) == 4
    # keep the listing the AI left out: re-applied right away
    assert api.review_override(run.name, {"id": dropped["id"], "keep": True})["reapplied"]
    assert len(api.products(run.name)["rows"]) == 5
    assert api.review_state(run.name)["overrides"] == 1
    api.review_override(run.name, {"id": dropped["id"], "keep": None})
    assert len(api.products(run.name)["rows"]) == 4
    audit = pd.read_csv(run / "ai_review" / "ai_classified_products.csv")
    assert "review_override" in audit.columns
    assert api.run_detail(run.name)["status"] == "reviewed"
    api.review_reset(run.name)
    assert not api.review_state(run.name)["prepared"]


def test_claude_review_runs_in_background(out, monkeypatch):
    run = make_run(out)
    api.review_prepare(run.name, {"instruction": "Keep shoes", "batch_size": 10})
    monkeypatch.setenv("ANTHROPIC_API_KEY", "test")
    monkeypatch.setattr(assist, "_call_claude", lambda prompt, **kw: "\n".join(
        f"p00{i + 1}|keep|0.9|shoe|ok" for i in range(5)))
    api.review_with_claude(run.name)
    assert api.wait_for(lambda: not api.review_state(run.name)["claude"]["running"])
    assert api.review_state(run.name)["counts"]["accepted"] == 5


@pytest.fixture()
def http(out):
    srv = server.make_server(0)
    thread = threading.Thread(target=srv.serve_forever, daemon=True)
    thread.start()
    yield f"http://127.0.0.1:{srv.server_address[1]}"
    srv.shutdown()
    srv.server_close()


def call(url, method="GET", body=None, headers=None):
    data = json.dumps(body).encode() if body is not None else None
    req = urllib.request.Request(url, data=data, method=method, headers={"Content-Type": "application/json", **(headers or {})})
    try:
        with urllib.request.urlopen(req) as res:
            return res.status, res.read(), dict(res.headers)
    except urllib.error.HTTPError as err:
        return err.code, err.read(), dict(err.headers)


def test_http_server_routes_and_guards(http, out):
    make_run(out)
    status, body, headers = call(http + "/")
    assert status == 200 and b"Price Lens" in body and "Content-Security-Policy" in headers
    assert call(http + "/js/app.js")[0] == 200
    assert call(http + "/../src/price_lens/webapp/api.py")[0] == 404
    status, body, _ = call(http + "/api/researches")
    assert status == 200 and json.loads(body)["items"][0]["listings"] == 5
    # changes need the app header (other websites cannot send it)
    assert call(http + "/api/drafts", "POST", {})[0] == 403
    status, body, _ = call(http + "/api/drafts", "POST", {}, {"X-PI-App": "1"})
    assert status == 200 and json.loads(body)["id"].startswith("draft_")
    # a foreign Host header (DNS rebinding) is refused
    assert call(http + "/api/meta", headers={"Host": "evil.example"})[0] == 403
    status, body, headers = call(http + "/api/runs/analyst_20260929_160000/export", "POST",
                                 {"kind": "csv", "rows": [0]}, {"X-PI-App": "1"})
    assert status == 200 and "attachment" in headers["Content-Disposition"]
    status, body, _ = call(http + "/api/runs/nope_1", headers={})
    assert status == 404 and json.loads(body)["error"]


ECB_XML = """<?xml version="1.0" encoding="UTF-8"?>
<gesmes:Envelope><Cube><Cube time='2026-09-28'>
<Cube currency='USD' rate='1.1500'/><Cube currency='GBP' rate='0.8500'/>
<Cube currency='PLN' rate='4.2500'/></Cube></Cube></gesmes:Envelope>"""


@pytest.fixture()
def ecb(monkeypatch):
    from price_lens.core import fx
    monkeypatch.setattr(fx, "_fetch_ecb", lambda timeout=10: ECB_XML)
    monkeypatch.delenv("PI_FX_OFFLINE", raising=False)
    api._fx_cache.clear()
    yield fx
    api._fx_cache.clear()


def test_exchange_rates_page_and_own_rates(out, ecb):
    fx = ecb
    data = api.exchange_rates()
    assert data["state"] == "live" and data["date"] == "2026-09-28"
    pln = next(r for r in data["rows"] if r["code"] == "PLN")
    assert pln["source"] == "ecb" and abs(pln["live"] - 1.15 / 4.25) < 1e-6 and pln["builtin"]
    aed = next(r for r in data["rows"] if r["code"] == "AED")
    assert aed["source"] == "builtin" and aed["live"] is None
    # "1 USD = 4 PLN", kept for future runs
    data = api.set_exchange_rate("pln", {"per_usd": 4, "keep": True, "note": "Client rate"})
    pln = next(r for r in data["rows"] if r["code"] == "PLN")
    assert pln["source"] == "yours" and pln["used"] == 0.25 and pln["diff_pct"] is not None
    # only for the next run
    api.set_exchange_rate("AED", {"per_usd": 3.5, "keep": False})
    assert next(r for r in api.exchange_rates()["rows"] if r["code"] == "AED")["source"] == "once"
    first = fx.rates_for_new_run(out)
    assert first["rates"]["PLN"] == 0.25 and abs(first["rates"]["AED"] - 1 / 3.5) < 1e-9
    assert set(first["manual"]) == {"PLN", "AED"} and "your own rates" in fx.describe(first)
    second = fx.rates_for_new_run(out)
    assert set(second["manual"]) == {"PLN"} and second["rates"]["AED"] == fx.STATIC_TO_USD["AED"]
    assert [h["action"] for h in api.exchange_rates()["history"]][:1] == ["used once"]
    api.remove_exchange_rate("PLN")
    assert "manual" not in fx.rates_for_new_run(out)
    with pytest.raises(api.ApiError):
        api.set_exchange_rate("PLN", {"per_usd": "abc"})
    with pytest.raises(api.ApiError):
        api.set_exchange_rate("PLN", {"per_usd": -2})
    with pytest.raises(api.ApiError):
        api.set_exchange_rate("PLNX", {"per_usd": 2})


def test_update_a_researchs_usd_prices(out, ecb):
    fx = ecb
    rows = [{"search_id": 1, "product_id": f"00{i}", "platform": "zalando", "marketplace": "zalando.pl",
             "product_name": f"Shoe {i}", "price_value": 400.0, "currency_code": "PLN", "price_usd": 100.0,
             "product_url": f"https://www.zalando.pl/p/{i}", "scraped_at": "2026-09-29T16:00:00"} for i in range(3)]
    rows.append({**rows[0], "product_id": "009", "marketplace": "zalando.de", "price_value": 100.0,
                 "currency_code": "EUR", "price_usd": 116.7})
    run = make_run(out, rows=rows)
    fx.save_for_run(run, fx.get_rates(out))
    api.set_exchange_rate("PLN", {"per_usd": 5, "keep": True, "note": "Client"})
    state = api.run_rates(run.name)
    pln = next(r for r in state["rows"] if r["code"] == "PLN")
    assert pln["listings"] == 3 and pln["changed"] and pln["now_source"] == "yours" and pln["was_source"] == "ecb"
    assert not next(r for r in state["rows"] if r["code"] == "EUR")["changed"]
    result = api.reprice_run(run.name, {"currencies": ["PLN"]})
    assert result["updated"] == ["PLN"]
    frame = pd.read_csv(run / "final_products.csv", dtype=str)
    assert list(frame.loc[frame.currency_code == "PLN", "price_usd"]) == ["80.00"] * 3
    assert frame.loc[frame.currency_code == "EUR", "price_usd"].iloc[0] == "116.7"  # untouched
    assert frame["product_id"].iloc[0] == "000"  # other columns written back as they were
    used = fx.load_for_run(run)
    assert used["rates"]["PLN"] == 0.2 and used["manual"]["PLN"]["note"] == "Client" and used["repriced"]
    assert "your rates for PLN" in api.run_detail(run.name)["fx"]
    assert not next(r for r in api.run_rates(run.name)["rows"] if r["code"] == "PLN")["changed"]
    with pytest.raises(api.ApiError):
        api.reprice_run(run.name, {"currencies": []})
    xlsx = api.export(run.name, {"kind": "xlsx", "rows": [0, 1, 2, 3]})
    from openpyxl import load_workbook
    import io
    about = [tuple(r) for r in load_workbook(io.BytesIO(xlsx.data))["About"].iter_rows(values_only=True)]
    assert any(r[0] == "Rates set by the analyst" and "PLN (Client)" in r[1] for r in about)


def test_retry_uses_the_researchs_rates(out, ecb, monkeypatch):
    from price_lens.orchestrator import jobs, retry
    parent = make_run(out)
    fx = ecb
    fx.save_for_run(parent, {**fx.get_rates(out), "rates": {**fx.get_rates(out)["rates"], "EUR": 2.0}})
    monkeypatch.setattr(jobs.subprocess, "Popen", lambda *a, **k: type("P", (), {"pid": 1})())
    child = jobs.start_job(PLAN, out, merge_into=parent)
    assert fx.load_for_run(child)["rates"]["EUR"] == 2.0
    pd.DataFrame([{"search_id": 1, "product_id": "n1", "platform": "zalando", "marketplace": "zalando.de",
                   "product_name": "New", "price_value": 10.0, "currency_code": "EUR", "price_usd": 11.67}]
                 ).to_csv(child / "detailed_products.csv", index=False)
    (child / "collection_report.json").write_text(json.dumps({"searches": []}), encoding="utf-8")
    retry.merge_retry(parent, child)
    merged = pd.read_csv(parent / "detailed_products.csv")
    assert merged.loc[merged.product_id == "n1", "price_usd"].iloc[0] == 20.0


def test_review_survives_a_moved_or_copied_project_folder(tmp_path, monkeypatch):
    """manifest.json stores absolute paths; a moved or copied folder must use its own files."""
    import shutil

    old_out = tmp_path / "old" / "output"
    monkeypatch.setenv("PI_OUTPUT_DIR", str(old_out))
    run = make_run(old_out)
    api.review_prepare(run.name, {"instruction": "Keep shoes", "batch_size": 3})
    prompt = api.review_prompt(run.name)
    first = prompt["batches"][:1]
    api.review_answer(run.name, {"batches": first, "text": "\n".join(
        f"p00{i + 1}|keep|0.95|shoe|ok" for i in range(3))})
    # the project is copied to a new folder; the old one stays (e.g. an older OneDrive copy)
    new_out = tmp_path / "new" / "output"
    shutil.copytree(old_out, new_out)
    monkeypatch.setenv("PI_OUTPUT_DIR", str(new_out))
    state = api.review_state(run.name)
    assert state["counts"]["accepted"] == 3 and state["counts"]["pending"] == 2
    rest = api.review_prompt(run.name)
    api.review_answer(run.name, {"batches": rest["batches"], "text": "p004|keep|0.9|shoe|ok\np005|drop|0.9|sock|no"})
    # the answer went into the new folder, not the old copy
    assert len(list((new_out / run.name / "ai_review" / "decisions").glob("*.json"))) == 2
    assert len(list((old_out / run.name / "ai_review" / "decisions").glob("*.json"))) == 1
    api.review_apply(run.name)
    assert len(api.products(run.name)["rows"]) == 4
    # and after the old folder is deleted, everything still works
    shutil.rmtree(tmp_path / "old")
    assert api.review_state(run.name)["counts"]["excluded"] == 1
    assert api.review_override(run.name, {"id": api.review_decisions(run.name)["rows"][0]["id"], "keep": False})
