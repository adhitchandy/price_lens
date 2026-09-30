import io
import json
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from price_lens.orchestrator import planning, price_report, storefront_status

ROOT = Path(__file__).parents[1]


def products():
    rows = []
    for market, currency, prices, usd_rate in (
        ("amazon.de", "EUR", [100, 200, 300], 1.167), ("zalando.de", "EUR", [90, 110], 1.167),
        ("ebay.co.uk", "GBP", [80, 120], 1.364),
    ):
        platform = market.split(".")[0]
        for i, price in enumerate(prices):
            rows.append({"search_id": 1, "platform": platform, "marketplace": market,
                         "product_name": f"Sneaker {i}", "brand": "Nike" if i else "Adidas",
                         "product_type": "Sneakers", "audience": "men" if i % 2 else "women",
                         "price_value": price, "currency_code": currency,
                         "price_usd": round(price * usd_rate, 2),
                         "product_url": f"https://{market}/p/{i}"})
    return pd.DataFrame(rows)


def test_prepare_summary_and_headline():
    df = price_report.prepare_products(products())
    assert set(df["country"]) == {"Germany", "United Kingdom"}
    summary = price_report.price_summary(df)
    assert {"country", "platform", "marketplace", "audience", "median_usd"} <= set(summary.columns)
    de_amazon_men = summary[(summary.marketplace == "amazon.de") & (summary.audience == "men")].iloc[0]
    assert de_amazon_men["listings"] == 1 and de_amazon_men["median"] == 200
    assert list(summary["median_usd"]) == sorted(summary["median_usd"])
    head = price_report.headline(df, summary)
    assert head["listings"] == 7 and head["storefronts"] == 3 and head["cheapest_country"] == "Germany"
    assert list(price_report.product_view(df).columns)[:5] == [
        "product_name", "product_url", "brand", "price_value", "currency_code"]


def test_price_report_workbook():
    df = price_report.prepare_products(products())
    data = price_report.price_report_xlsx(price_report.price_summary(df), df,
                                          {"research_question": "Sneakers", "run": "x"})
    book = load_workbook(io.BytesIO(data))
    assert book.sheetnames == ["Summary", "Products", "About"]
    assert book["Summary"]["A1"].value == "Country"
    assert book["Products"]["A1"].value == "Product" and book["Products"].max_row == 8


def test_empty_run_does_not_break_summary():
    df = price_report.prepare_products(pd.DataFrame())
    assert price_report.price_summary(df).empty


def search(targets, translations=None):
    return {"query": {"default": "sneakers", "translations": translations or {}, "filters": {}},
            "targets": targets}


def test_country_groups_and_presets():
    s = search([{"platform": "amazon", "countries": []}, {"platform": "zalando", "countries": []}])
    notes = planning.apply_country_group(s, "DACH")
    assert s["targets"][0]["countries"] == ["Germany"]           # no Amazon AT/CH
    assert s["targets"][1]["countries"] == ["Germany", "Austria", "Switzerland"]
    assert any("Amazon" in n for n in notes)
    notes = planning.apply_product_preset(s, "Electronics & appliances")
    assert [t["platform"] for t in s["targets"]] == ["mediamarkt", "amazon", "ebay"]
    assert s["targets"][0]["countries"] == ["Germany", "Austria", "Switzerland"]
    assert "screen protector" in s["query"]["filters"]["exclude"] and s["audiences"] == []


def test_missing_translations_and_estimate():
    s = search([{"platform": "mediamarkt", "countries": ["Germany", "Turkey", "Austria"]}],
               {"de": "Sneaker"})
    assert planning.missing_translations(s) == {"tr": ["Turkey"]}
    preview = [{"platform": "mediamarkt", "domain": "mediamarkt.de", "audience": ""},
               {"platform": "zalando", "domain": "zalando.de", "audience": "men, women"}]
    small = planning.estimate_minutes(preview, {"products_per_storefront": 10})
    big = planning.estimate_minutes(preview, {"products_per_storefront": 500})
    assert 1 <= small < big


def test_storefront_status_memory(tmp_path):
    reports = [{"search_id": 1, "status": "partial", "report": {"platform_reports": {"mediamarkt": {
        "events": [{"marketplace": "mediamarkt.de", "listings": 12, "status": "succeeded"}],
        "failures": [{"marketplace": "mediamarkt.be",
                      "error": "Access restriction or verification challenge; stopped"}]}}}}]
    preview = [{"platform": "mediamarkt", "domain": d, "country": c}
               for d, c in (("mediamarkt.de", "Germany"), ("mediamarkt.be", "Belgium"),
                            ("mediamarkt.at", "Austria"))]
    storefront_status.record(tmp_path, reports, preview, "run1")
    status = storefront_status.load(tmp_path)
    assert status["mediamarkt.de"]["result"] == "OK"
    assert status["mediamarkt.be"]["result"] == "Bot check"
    assert "mediamarkt.at" not in status  # never reached: no evidence recorded
    assert list(storefront_status.problems(status, ["mediamarkt.de", "mediamarkt.be"])) == ["mediamarkt.be"]


def component_data(app, name):
    """The JSON a custom screen (Streamlit components v2) received in the test harness."""
    found = []

    def walk(node):
        children = getattr(node, "children", None)
        for child in (children.values() if isinstance(children, dict) else children or []):
            if getattr(child, "type", "") == "bidi_component" and child.proto.component_name == name:
                found.append(json.loads(child.proto.json))
            walk(child)
    walk(app._tree)
    assert found, f"component {name} not rendered"
    return found[-1]


def open_run(tmp_path, monkeypatch):
    from streamlit.testing.v1 import AppTest

    monkeypatch.setenv("PI_OUTPUT_DIR", str(tmp_path))
    app = AppTest.from_file(str(ROOT / "apps/unified_app.py"), default_timeout=60).run()
    next(b for b in app.button if b.label == "📂 Open in Results").click().run()
    assert not app.exception
    return app


def test_results_tab_renders_price_research(tmp_path, monkeypatch):
    run = tmp_path / "analyst_20260929_150000"
    run.mkdir()
    products().to_csv(run / "final_products.csv", index=False)
    (run / "collection_report.json").write_text(json.dumps({"searches": [
        {"search_id": 1, "status": "completed", "report": {}}]}))
    (run / "analyst_plan.json").write_text(json.dumps({"research_question": "Sneakers", "searches": [{}]}))
    data = component_data(open_run(tmp_path, monkeypatch), "pi_results")
    assert data["title"] == "Sneakers" and data["source"] == "all" and not data["has_review"]
    assert len(data["rows"]) == 7 and data["rows"][0]["i"] == 0
    first = data["rows"][0]
    assert first["c"] == "Germany" and first["m"] == "amazon.de" and first["usd"] == 116.7
    assert data["fx"] == "USD at static rates"


def test_results_compare_one_product_group_and_flag_stale_review(tmp_path, monkeypatch):
    from price_lens.agent.review import mark_review_outdated, prepare_review_batches

    run = tmp_path / "analyst_20260929_160000"
    run.mkdir()
    frame = products()
    frame.loc[frame.index[-2:], "search_id"] = 2  # the eBay rows are a different product
    frame["relevance_confidence_pct"] = [95, 90, 40, 80, 70, 60, 99]
    frame.to_csv(run / "final_products.csv", index=False)
    frame.to_csv(run / "detailed_products.csv", index=False)
    (run / "collection_report.json").write_text(json.dumps({"searches": [
        {"search_id": 1, "status": "completed", "report": {}},
        {"search_id": 2, "status": "completed", "report": {}}]}))
    (run / "analyst_plan.json").write_text(json.dumps({"research_question": "Clothes", "searches": [
        {"product_type": "Sneakers"}, {"product_type": "Hoodies"}]}))
    prepare_review_batches(run / "detailed_products.csv", run / "ai_review",
                           research_question="Clothes", instruction="Keep", batch_size=10)
    frame.head(3).to_csv(run / "ai_review" / "final_products.csv", index=False)
    mark_review_outdated(run / "ai_review", "retry retry_x", 3)
    app = open_run(tmp_path, monkeypatch)
    data = component_data(app, "pi_results")
    assert data["outdated"]["added"] == 3 and data["has_review"]
    assert data["source"] == "all" and len(data["rows"]) == 7  # stale review: everything
    assert data["groups"] == ["1 · Sneakers", "2 · Hoodies"]
    assert {r["g"] for r in data["rows"]} == {"1 · Sneakers", "2 · Hoodies"}
    assert [r["conf"] for r in data["rows"]][:3] == [95, 90, 40]
    assert any(b.label.startswith("➕ Add the 3 new") for b in app.button)


def test_results_export_uses_the_rows_on_screen(tmp_path, monkeypatch):
    from apps.pi_ui import components, results_tab
    from price_lens.orchestrator import price_report

    run = tmp_path / "analyst_20260929_170000"
    run.mkdir()
    prepared = price_report.prepare_products(products()).reset_index(drop=True)
    result = {"final": products(), "run_dir": str(run)}
    request = {"kind": "xlsx", "rows": [0, 3, 4, 99], "filters": "country=Germany", "group": ""}
    exported = components.export_file(prepared, request, {"run": run.name}, None)
    book = load_workbook(io.BytesIO(__import__("base64").b64decode(exported["b64"])))
    assert exported["rows"] == 3 and book["Products"].max_row == 4  # header + 3 rows
    about = {r[0].value: r[1].value for r in book["About"].iter_rows()}
    assert about["Filters applied in app"] == "country=Germany"

    state = {"pi_results_x": {"export": {**request, "kind": "csv"}, "source": "all"}}
    monkeypatch.setattr(results_tab.st, "session_state", state)
    results_tab._on_export("pi_results_x", run, result, {}, {"run": run.name}, None)
    reply = state["pi_results_x__export"]
    csv = __import__("base64").b64decode(reply["b64"]).decode("utf-8-sig")
    assert reply["name"] == f"products_{run.name}.csv" and len(csv.strip().splitlines()) == 4
    assert csv.splitlines()[0].startswith("Product,Link,Brand,Price")

    state["pi_results_x"]["export"] = {"kind": "xlsx", "rows": "not a list"}
    results_tab._on_export("pi_results_x", run, result, {}, {"run": run.name}, None)
    assert "error" in state["pi_results_x__export"]
