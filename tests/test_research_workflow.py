import json

import pandas as pd
import pytest
from openpyxl import load_workbook

from price_lens.agent.review import apply_review_decisions, prepare_review_batches
from price_lens.core.results import ScrapeOutcome
from price_lens.core.schemas import ScrapeReport
from price_lens.core.validation import RequestValidationError
from price_lens.orchestrator.registry import MarketplaceAdapter
from price_lens.orchestrator.runner import run_research
from price_lens.orchestrator.validation import parse_research_plan


def _plan(tmp_path):
    return parse_research_plan(
        {
            "research_question": "Compare men's running shoes in Germany",
            "marketplaces": {
                "amazon": ["amazon.de"],
                "ebay": ["ebay.de"],
                "zalando": ["zalando.de"],
            },
            "searches": [
                {
                    "search_term": "black running shoes",
                    "product_type": "Running Shoes",
                    "audiences": ["men"],
                    "exclude": ["kids", "used"],
                }
            ],
            "pages": 1,
            "output_dir": str(tmp_path),
            "cleaning": {"remove_outliers": False},
            "review": {"instruction": "Keep men's running shoes", "batch_size": 5},
        }
    )


def test_research_plan_compiles_for_all_selected_platforms(tmp_path):
    plan = _plan(tmp_path)

    assert plan.platforms == ("amazon", "ebay", "zalando")
    assert plan.platform_plan("amazon").searches[0].audiences == ()
    assert plan.platform_plan("zalando").searches[0].audiences == ("men",)


def test_research_plan_requires_audience_when_searching_zalando(tmp_path):
    with pytest.raises(RequestValidationError, match="requires 'audiences'"):
        parse_research_plan(
            {
                "research_question": "Find shoes",
                "marketplaces": {"zalando": ["zalando.de"]},
                "searches": [
                    {"search_term": "shoes", "product_type": "Shoes"}
                ],
                "output_dir": str(tmp_path),
            }
        )


def test_research_plan_accepts_mediamarkt_without_audience(tmp_path):
    plan = parse_research_plan(
        {
            "research_question": "Compare Pixel prices",
            "marketplaces": {"mediamarkt": ["mediamarkt.de"]},
            "searches": [{"search_term": "Google Pixel", "product_type": "Smartphone"}],
            "output_dir": str(tmp_path),
        }
    )

    assert plan.platforms == ("mediamarkt",)
    assert plan.platform_plan("mediamarkt").searches[0].audiences == ()


def test_unified_runner_writes_platform_and_combined_outputs(tmp_path):
    plan = _plan(tmp_path)

    def adapter(platform):
        def collect(_plan, _log):
            frame = pd.DataFrame(
                [
                    {
                        "platform": platform,
                        "marketplace": f"{platform}.de",
                        "product_name": f"{platform} shoe",
                        "origin_country": "Germany",
                        "price_usd": 50.0,
                        "price_value": 45.0,
                        "currency_code": "EUR",
                        "product_id": f"{platform}-1",
                        "product_url": f"https://example.com/{platform}-1",
                        "search_term": "black running shoes",
                        "product_type": "Running Shoes",
                        "sponsored": False,
                        "rating": 4.5,
                    }
                ]
            )
            report = ScrapeReport(platform=platform, status="completed")
            return ScrapeOutcome(frame, report)

        return MarketplaceAdapter(platform, collect)

    adapters = {platform: adapter(platform) for platform in plan.platforms}
    result = run_research(plan, adapters=adapters, log=lambda _message: None)

    assert len(result.cleaned_products) == 3
    assert (result.run_dir / "combined_cleaned_products.csv").exists()
    assert (result.run_dir / "amazon" / "cleaned_products.csv").exists()
    assert result.report["status"] == "completed"


def test_review_batches_are_validated_and_merged(tmp_path):
    source = tmp_path / "products.csv"
    pd.DataFrame(
        [
            {"platform": "amazon", "marketplace": "amazon.de", "product_name": "Shoe"},
            {"platform": "ebay", "marketplace": "ebay.de", "product_name": "Shoe bag"},
        ]
    ).to_csv(source, index=False)
    review_dir = tmp_path / "review"
    manifest = prepare_review_batches(
        source,
        review_dir,
        research_question="Find running shoes",
        instruction="Exclude accessories",
        batch_size=5,
        target_results=1,
    )
    batch = json.loads((review_dir / "batch_001.json").read_text(encoding="utf-8"))
    decisions = {
        "decisions": [
            {
                "row_id": product["row_id"],
                "keep": index == 0,
                "category": "shoe" if index == 0 else "accessory",
                "confidence": 0.95,
                "reason": "Complete shoe" if index == 0 else "Bag only",
            }
            for index, product in enumerate(batch["products"])
        ]
    }
    decision_path = manifest["batches"][0]["decisions"]
    with open(decision_path, "w", encoding="utf-8") as handle:
        json.dump(decisions, handle)

    report = apply_review_decisions(review_dir)

    assert report["reviewed"] == 2
    assert report["accepted"] == 1
    assert report["candidates"] == 1
    coverage = json.loads((review_dir / "coverage_report.json").read_text(encoding="utf-8"))
    assert coverage["target_met_by_accepted"] is True
    assert coverage["by_platform"] == {"amazon": 1}
    final = pd.read_csv(review_dir / "final_products.csv")
    assert final.iloc[0]["product_name"] == "Shoe"
    assert final.iloc[0]["review_status"] == "accepted"
    assert final.iloc[0]["relevance_confidence_pct"] == 95.0
    assert "ai_reason" not in final.columns
    assert "price_text" not in final.columns
    audit = pd.read_csv(review_dir / "ai_classified_products.csv")
    assert len(audit) == 2
    assert "ai_reason" in audit.columns
    workbook = load_workbook(review_dir / "final_products.xlsx", read_only=False)
    worksheet = workbook["Products"]
    headers = [cell.value for cell in worksheet[1]]
    assert headers == list(final.columns)
    assert worksheet.freeze_panes == "E2"
    assert "FinalProducts" in worksheet.tables
    workbook.close()


def test_unified_runner_continues_when_one_platform_cannot_start(tmp_path):
    plan = _plan(tmp_path)

    def failed(_plan, _log):
        raise RuntimeError("browser unavailable")

    def succeeded(_plan, _log):
        frame = pd.DataFrame(
            [
                {
                    "platform": "ebay",
                    "marketplace": "ebay.de",
                    "product_name": "Running shoe",
                    "origin_country": "Germany",
                    "price_usd": 50.0,
                    "price_value": 45.0,
                    "currency_code": "EUR",
                    "product_id": "1",
                    "product_url": "https://example.com/1",
                    "search_term": "shoes",
                    "product_type": "Running Shoes",
                    "sponsored": False,
                    "rating": 4.5,
                }
            ]
        )
        return ScrapeOutcome(frame, ScrapeReport(platform="ebay", status="completed"))

    adapters = {
        "amazon": MarketplaceAdapter("amazon", failed),
        "ebay": MarketplaceAdapter("ebay", succeeded),
        "zalando": MarketplaceAdapter("zalando", failed),
    }
    result = run_research(plan, adapters=adapters, log=lambda _message: None)

    assert result.report["status"] == "partial"
    assert result.report["failed_platforms"] == ["amazon", "zalando"]
    assert len(result.cleaned_products) == 1


def test_unified_runner_restarts_marketplace_once_after_browser_failure(tmp_path):
    plan = parse_research_plan(
        {
            "research_question": "Find shoes",
            "marketplaces": {"amazon": ["amazon.de"]},
            "searches": [{"search_term": "shoes", "product_type": "Shoes"}],
            "output_dir": str(tmp_path),
        }
    )
    calls = 0

    def flaky(_plan, _log):
        nonlocal calls
        calls += 1
        if calls == 1:
            raise RuntimeError("browser exited")
        frame = pd.DataFrame(
            [
                {
                    "platform": "amazon",
                    "product_name": "Shoe",
                    "origin_country": "Germany",
                    "price_value": 50.0,
                    "price_usd": 50.0,
                    "sponsored": False,
                    "product_id": "1",
                    "rating": 4.0,
                }
            ]
        )
        return ScrapeOutcome(frame, ScrapeReport(platform="amazon", status="completed"))

    result = run_research(
        plan,
        adapters={"amazon": MarketplaceAdapter("amazon", flaky)},
        log=lambda _message: None,
    )

    assert calls == 2
    assert result.report["status"] == "completed"
    assert result.report["platform_reports"]["amazon"]["collection_attempts"] == 2


def test_review_quality_gates_keep_uncertain_candidates_in_one_final_file(tmp_path):
    source = tmp_path / "products.csv"
    pd.DataFrame(
        [
            {"platform": "amazon", "product_name": "Nike Pegasus", "brand": "Nike"},
            {"platform": "amazon", "product_name": "Running shoe", "brand": "Altra"},
            {"platform": "amazon", "product_name": "Adidas Duramo", "brand": "Adidas"},
            {"platform": "amazon", "product_name": "Nike unknown model", "brand": "Nike"},
        ]
    ).to_csv(source, index=False)
    review_dir = tmp_path / "review"
    manifest = prepare_review_batches(
        source,
        review_dir,
        research_question="Nike or Adidas shoes",
        instruction="Keep target shoes",
        batch_size=5,
        minimum_keep_confidence=0.8,
        allowed_brands=("nike", "adidas"),
    )
    batch = json.loads((review_dir / "batch_001.json").read_text(encoding="utf-8"))
    confidences = [0.95, 0.99, 0.7, 0.6]
    decisions = {
        "decisions": [
            {
                "row_id": product["row_id"],
                "keep": index != 3,
                "category": "running shoes",
                "confidence": confidences[index],
                "reason": "Agent says keep",
            }
            for index, product in enumerate(batch["products"])
        ]
    }
    with open(manifest["batches"][0]["decisions"], "w", encoding="utf-8") as handle:
        json.dump(decisions, handle)

    report = apply_review_decisions(review_dir)

    assert report["accepted"] == 1
    assert report["candidates"] == 3
    assert report["excluded"] == 1
    assert report["uncertain"] == 2
    final = pd.read_csv(review_dir / "final_products.csv")
    assert set(final["review_status"]) == {"accepted", "uncertain"}
    assert set(final["product_name"]) == {
        "Nike Pegasus",
        "Adidas Duramo",
        "Nike unknown model",
    }
    relevance = dict(zip(final["product_name"], final["relevance_confidence_pct"]))
    assert relevance["Nike Pegasus"] == 95.0
    assert relevance["Adidas Duramo"] == 70.0
    assert relevance["Nike unknown model"] == 40.0
    assert not (review_dir / "needs_manual_review.csv").exists()
    assert not (review_dir / "ai_exclusions.csv").exists()
    audit = pd.read_csv(review_dir / "ai_classified_products.csv")
    statuses = dict(zip(audit["product_name"], audit["review_status"]))
    assert statuses["Running shoe"] == "excluded"
